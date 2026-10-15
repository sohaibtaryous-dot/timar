# -*- coding: utf-8 -*-
"""
Module EXPORT — Extraction BL / Packing Lists + Factures commerciales vers 2 fichiers Excel :

  1. FACTURES.xlsx                → feuille unique « Factures » (5 colonnes, import système strict)
  2. BL_REFERENCES_SUMMARY.xlsx   → « BL References », « BL SUMMARY », « GLOBAL TOTAL »

Règle prioritaire : une ligne d'article source = une ligne Excel (jamais de regroupement).
"""

from __future__ import annotations

import io
import re
import zipfile
from collections import Counter, OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import app as core  # moteur commun (lecture PDF, OCR, tableaux, nombres)

FACTURES_COLUMNS = ["Numéro de facture", "Nom du fournisseur ou client", "Date de la facture",
                    "Devise", "Montant en devise"]
BLREF_COLUMNS = ["Reference", "Description", "Quantity", "Unit", "BL Number (Bon de Livraison)",
                 "Invoice Number (Facture)", "Total Net Weight of the BL", "Total Gross Weight of the BL",
                 "Number of Packages (Nombre de Colis)", "Prix unitaire total (Hors-taxe)", "Customer PO"]
SUMMARY_COLUMNS = ["BL Number", "Total Quantity", "Total Net Weight", "Total Gross Weight",
                   "Number of Packages", "Total Valeur HT"]
GLOBAL_FIELDS = ["COLIS", "POIDS BRUT", "POIDS NET", "VALEUR DEVISE", "QUANTITE"]

# =============================================================================
# DATES
# =============================================================================

MONTHS = {
    "jan": 1, "janv": 1, "january": 1, "janvier": 1, "feb": 2, "fev": 2, "fevr": 2, "february": 2, "fevrier": 2,
    "mar": 3, "mars": 3, "march": 3, "apr": 4, "avr": 4, "april": 4, "avril": 4, "may": 5, "mai": 5,
    "jun": 6, "juin": 6, "june": 6, "jul": 7, "juil": 7, "july": 7, "juillet": 7, "aug": 8, "aou": 8,
    "aout": 8, "august": 8, "sep": 9, "sept": 9, "september": 9, "septembre": 9, "oct": 10, "october": 10,
    "octobre": 10, "nov": 11, "november": 11, "novembre": 11, "dec": 12, "december": 12, "decembre": 12,
}
DATE_NUM_RE = re.compile(r"\b(\d{1,2})[./-](\d{1,2})[./-](\d{4}|\d{2})\b|\b(\d{4})[./-](\d{1,2})[./-](\d{1,2})\b")
DATE_TXT_RE = re.compile(r"\b(\d{1,2})(?:er)?[\s./-]*([A-Za-zÀ-ÿ]{3,9})\.?[\s./,-]*(\d{4}|\d{2})\b")
DATE_TXT_EN_RE = re.compile(r"\b([A-Za-z]{3,9})\.?\s+(\d{1,2}),?\s+(\d{4})\b")


def _year(y: str) -> int:
    y = int(y)
    return y + 2000 if y < 100 else y


def _mk(y, m, d):
    try:
        return date(_year(str(y)), int(m), int(d))
    except (ValueError, TypeError):
        return None


def parse_date_any(text: str):
    """Première date valide trouvée dans `text` (jour en premier : format européen)."""
    if not text:
        return None
    cands = []
    for m in DATE_NUM_RE.finditer(text):
        dt = _mk(m.group(3), m.group(2), m.group(1)) if m.group(1) else _mk(m.group(4), m.group(5), m.group(6))
        if dt:
            cands.append((m.start(), dt))
    for m in DATE_TXT_RE.finditer(text):
        mon = MONTHS.get(core.strip_accents(m.group(2)).lower().rstrip("."))
        if mon:
            dt = _mk(m.group(3), mon, m.group(1))
            if dt:
                cands.append((m.start(), dt))
    for m in DATE_TXT_EN_RE.finditer(text):
        mon = MONTHS.get(core.strip_accents(m.group(1)).lower())
        if mon:
            dt = _mk(m.group(3), mon, m.group(2))
            if dt:
                cands.append((m.start(), dt))
    return min(cands)[1] if cands else None


def find_invoice_date(text: str):
    lines = core.text_lines(text)
    for i, ln in enumerate(lines[:80]):
        if re.search(r"(invoice\s*date|date\s*(de\s*(la\s*)?)?facture|date\s*of\s*invoice|rechnungsdatum|fecha\s*factura|^date\b|\bdate\s*:)",
                     ln, re.I) and not re.search(r"(due|échéance|echeance|delivery|livraison|order|commande|shipping|expédition)",
                                                 ln, re.I):
            dt = parse_date_any(ln) or (parse_date_any(lines[i + 1]) if i + 1 < len(lines) else None)
            if dt:
                return dt
    return parse_date_any("\n".join(lines[:40]))


# =============================================================================
# EN-TÊTES EXPORT
# =============================================================================

CLIENT_LABEL_RE = re.compile(
    r"^\s*(?:sold\s*to|bill(?:ed)?\s*to|invoice\s*to|invoiced\s*to|factur[ée]e?\s*[àa]|adresse\s*de\s*facturation|"
    r"client|customer|buyer|acheteur|destinataire\s*de\s*la\s*facture|rechnungsempf[äa]nger|kunde)"
    r"(?!\s*(?:p\.?\s?o|code|n°|no\b|nr|number|num|ref|order|part|id|vat|tva))\s*[:\-]?\s*(.*)$", re.I)
SPLIT_COLS_RE = re.compile(r"\s{2,}|\s+(?:ship\s*to|deliver\w*\s*to|livr[ée]\w*\s*[àa]|consignee|destinataire|"
                           r"invoice\s*(?:no|n°|number|date)|date)\b", re.I)
PO_RE = re.compile(
    r"(?:customer\s*p\.?\s?o\.?|client\s*p\.?\s?o\.?|\bp\.?\s?o\.?\s*(?:no\.?|n[°º]|nr\.?|number|#)?|purchase\s*order\s*(?:no\.?|n[°º]|number|#)?|"
    r"(?:your|votre)\s*(?:order|commande)\s*(?:no\.?|n[°º]|number|#)?|commande\s*client\s*(?:n[°º])?|"
    r"order\s*(?:no\.?|n[°º]|number|#))\s*[:#.]?\s*([A-Z0-9][A-Z0-9\-/_.]{3,})", re.I)
BL_EXPORT_RE = re.compile(
    r"(?:delivery\s*note|delivery|bon\s*de\s*livraison|livraison|\bb\.?l\.?\b|lieferschein|packing\s*list|"
    r"liste\s*de\s*colisage|despatch\s*note|dispatch\s*note|shipment)\s*"
    r"(?:no\.?|n[°º]\.?|nr\.?|number|num[ée]ro|#)\s*[:.#]?\s*([A-Z0-9][A-Z0-9\-/_]*\d[A-Z0-9\-/_]*)", re.I)
HT_PATTERNS = [
    r"total\s*(?:amount\s*)?(?:\(?h\.?t\.?\)?\b|hors\s*tax\w*|excl\w*\.?\s*(?:of\s*)?(?:vat|tax\w*)|net\s*of\s*tax|before\s*tax|w/o\s*vat)",
    r"montant\s*(?:total\s*)?(?:\(?h\.?t\.?\)?|hors\s*tax\w*)",
    r"total\s*net\s*(?:amount|value|invoice|ht)?(?!\s*weight)(?!\s*wt)",
    r"net\s*(?:amount|value|invoice\s*value)",
    r"sub\s*-?\s*total|sous[- ]?total|zwischensumme|nettobetrag|netto\s*summe",
    r"total\s*(?:invoice\s*)?(?:value|amount|goods|marchandises)",
    r"(?:invoice\s*)?total\b",
]
NET_A_PAYER_RE = re.compile(r"(net\s*[àa]\s*payer|amount\s*due|total\s*due|balance\s*due)", re.I)
TTC_RE = re.compile(r"(\bttc\b|incl\w*\.?\s*(?:vat|tax)|total\s*(?:amount\s*)?(?:with|including)\s*vat|brutto)", re.I)
VAT_ZERO_RE = re.compile(r"(?:tva|vat|mwst|iva)[^\n]{0,25}?\b0(?:[.,]0+)?\s*%|exon[ée]r|exempt|0\s*%\s*(?:tva|vat)", re.I)
NOT_AMOUNT_LINE = re.compile(r"(weight|poids|kg\b|kgs|gewicht|colis|packages|pallets|qty|quantit|pieces|pcs)", re.I)


def _line_amount(line: str, dec: str):
    vals = []
    for tok in line.replace("€", " ").split():
        v = core.parse_number(tok.strip(":;"), dec)
        if v is not None and re.search(r"\d", tok) and not core.DATE_RE.fullmatch(tok.strip(":;")):
            vals.append(v)
    return vals[-1] if vals else None


def find_total_ht(text: str, dec: str, lines_sum: float | None, notes: list, inv_label: str):
    lines = core.text_lines(text)
    for pat in HT_PATTERNS:
        rx = re.compile(pat, re.I)
        for i, ln in enumerate(lines):
            if not rx.search(ln) or TTC_RE.search(ln) or NOT_AMOUNT_LINE.search(ln):
                continue
            v = _line_amount(ln[rx.search(ln).end():], dec)
            if v is None and i + 1 < len(lines) and not re.search(r"[A-Za-z]{4}", lines[i + 1]):
                v = _line_amount(lines[i + 1], dec)
            if v is not None and v > 0:
                return round(v, 2), "Total HT"
    # Net à payer accepté uniquement si TVA 0 % / exonéré et pas de TTC distinct
    for ln in lines:
        if NET_A_PAYER_RE.search(ln):
            v = _line_amount(ln, dec)
            if v is not None and (VAT_ZERO_RE.search(text) or not TTC_RE.search(text)):
                notes.append(f"Facture {inv_label} : montant pris sur « Net à payer » (TVA 0 % / pas de TTC distinct).")
                return round(v, 2), "Net à payer"
    if lines_sum:
        notes.append(f"Facture {inv_label} : total HT non trouvé, montant = somme des lignes ({lines_sum:.2f}).")
        return round(lines_sum, 2), "Somme des lignes"
    return None, ""


def find_client(text: str) -> str:
    lines = core.text_lines(text)
    for i, ln in enumerate(lines[:60]):
        m = CLIENT_LABEL_RE.match(ln)
        if not m:
            continue
        cand = SPLIT_COLS_RE.split(m.group(1).strip(" :-"))[0].strip(" ,;:-")
        if not re.search(r"[A-Za-z]{3}", cand or ""):
            for nxt in lines[i + 1:i + 3]:
                c2 = SPLIT_COLS_RE.split(nxt)[0].strip(" ,;:-")
                if re.search(r"[A-Za-z]{3}", c2):
                    cand = c2
                    break
        if re.search(r"[A-Za-z]{3}", cand or ""):
            return cand[:100]
    return ""


def find_po(text: str) -> str:
    for m in PO_RE.finditer(text):
        val = m.group(1).strip(".-/")
        if re.search(r"\d", val) and not core.DATE_RE.fullmatch(val):
            return val
    return ""


def find_bl_numbers(text: str) -> list[str]:
    out = []
    for rx in (BL_EXPORT_RE, core.BL_NO_RE, core.BL_NO_RE2):
        for m in rx.finditer(text):
            val = m.group(1).strip(".-/")
            if len(val) >= 3 and not core.DATE_RE.fullmatch(val) and val not in out:
                out.append(val)
    return out


# =============================================================================
# DOCUMENTS
# =============================================================================

@dataclass
class ExpInvoice:
    doc: object
    number: str = ""
    name: str = ""
    supplier: str = ""
    date: object = None
    currency: str = ""
    total_ht: float | None = None
    total_src: str = ""
    po: str = ""
    bl_refs: list = field(default_factory=list)
    lines: list = field(default_factory=list)   # articles
    fees: list = field(default_factory=list)    # frais (fret, emballage…)
    net: float | None = None
    gross: float | None = None
    packages: int | None = None


@dataclass
class ExpBL:
    doc: object
    number: str = ""
    po: str = ""
    invoice_refs: list = field(default_factory=list)
    net: float | None = None
    gross: float | None = None
    packages: int | None = None
    lines: list = field(default_factory=list)   # {pos, ref, desc, qty, unit, po}
    invoices: list = field(default_factory=list)


UNIT_ALT = "|".join(sorted(core.UNITS, key=len, reverse=True))
BL_LINE_RE = re.compile(rf"^(?:(?P<pos>\d{{1,5}})[.)]?\s+)?(?P<ref>[A-Z0-9][A-Z0-9\-_./]*\d[A-Z0-9\-_./]*)\s+(?P<desc>.*?)\s+"
                        rf"(?:(?P<qty>\d[\d.,]*)\s*(?P<unit>{UNIT_ALT})\b|(?P<unit2>{UNIT_ALT})\s+(?P<qty2>\d[\d.,]*))"
                        rf"(?P<rest>.*)$", re.I)


def clean_description(parts: list[str]) -> str:
    keep = []
    for p in parts:
        for seg in str(p).split("\n"):
            seg = seg.strip()
            if not seg:
                continue
            if core.HS_LABEL_RE.search(seg) or core.LINE_NET_RE.search(seg) or core.LINE_GROSS_RE.search(seg) \
                    or PO_RE.search(seg) or re.search(r"(serial\s*n|n°\s*de\s*s[ée]rie|s/n\s*:)", seg, re.I):
                continue
            keep.append(seg)
    return re.sub(r"\s+", " ", " ".join(keep)).strip()


def _qty(v):
    if v is None:
        return None
    return int(v) if float(v).is_integer() else float(v)


def parse_bl_lines(doc, dec) -> list[dict]:
    out = []
    for page in doc.pages:
        found = []
        for tset in page.table_sets():
            found = []
            for header, row in core.iter_table_rows(
                    tset, core.LINE_COLMAP, lambda m: ("ref" in m or "alstom" in m) and "qty" in m):
                ref = core.row_cell(row, header, "ref").split("\n")[0] or core.row_cell(row, header, "alstom").split("\n")[0]
                qtxt = core.row_cell(row, header, "qty")
                qty = core.parse_number(qtxt.split()[0], dec) if qtxt else None
                if not ref or qty is None or core.TOTAL_WORDS.search(ref):
                    continue
                unit = core.row_cell(row, header, "unit")
                if not unit:
                    um = re.search(rf"\b({UNIT_ALT})\b", qtxt, re.I)
                    unit = um.group(1) if um else ""
                found.append({"pos": core.row_cell(row, header, "pos"), "ref": ref,
                              "desc": clean_description([core.row_cell(row, header, "desc")]),
                              "qty": _qty(qty), "unit": unit.strip(), "po": core.row_cell(row, header, "po")})
            if found:
                break
        if not found:  # analyse ligne à ligne
            last = None
            for ln in core.text_lines(page.text):
                m = BL_LINE_RE.match(ln)
                if m and not core.TOTAL_WORDS.search(ln):
                    qty = core.parse_number(m.group("qty") or m.group("qty2"), dec)
                    if qty is None:
                        continue
                    last = {"pos": m.group("pos") or "", "ref": m.group("ref"), "desc": clean_description([m.group("desc")]),
                            "qty": _qty(qty), "unit": (m.group("unit") or m.group("unit2") or "").strip(), "po": ""}
                    found.append(last)
                elif last is not None and len(ln) < 90 and not core.TOTAL_WORDS.search(ln) \
                        and sum(c.isalpha() for c in ln) >= 0.6 * len(ln.replace(" ", "")) \
                        and not re.search(r"(weight|poids|colis|packages|page)", ln, re.I):
                    extra = clean_description([ln])
                    if extra:
                        last["desc"] = f"{last['desc']} {extra}".strip()
                else:
                    last = None if core.TOTAL_WORDS.search(ln) else last
        out.extend(found)
    return out


def parse_export_documents(docs, client_mode: str, notes: list):
    invoices, bls = [], []
    for d in docs:
        text = d.text
        dec = core.detect_decimal(text)
        if d.kind == "FACTURE":
            base = core.parse_invoice(d)
            inv = ExpInvoice(doc=d, number=base.number, supplier=base.supplier, currency=base.currency,
                             date=find_invoice_date(text), po=find_po(text), bl_refs=find_bl_numbers(text),
                             net=base.net_total, gross=base.gross_total, packages=base.packages)
            inv.name = (find_client(text) if client_mode == "client" else base.supplier) or ""
            for it in base.items:
                desc = clean_description([it.get("desc_full") or it.get("desc", "")] + it.get("desc_extra", []))
                line = {"pos": core.clean_cell(it.get("pos")), "ref": core.clean_cell(it.get("ref")) or core.clean_cell(it.get("alstom")),
                        "desc": desc, "qty": _qty(it.get("qty")),
                        "unit": core.clean_cell(it.get("unit_raw")) or core.clean_cell(it.get("unit")),
                        "amount": round(it["amount"], 2) if it.get("amount") is not None else None,
                        "po": core.clean_cell(it.get("po")), "_used": False}
                if core.FEE_WORDS.search(f"{desc} {line['ref']}") and not re.search(r"\d", line["ref"] or ""):
                    inv.fees.append(line)
                else:
                    inv.lines.append(line)
            lines_sum = sum(l["amount"] for l in inv.lines + inv.fees if l["amount"] is not None) or None
            label = inv.number or f"({d.file} p.{d.page_label})"
            inv.total_ht, inv.total_src = find_total_ht(text, dec, lines_sum, notes, label)
            invoices.append(inv)
        elif d.kind == "PACKING":
            nums = find_bl_numbers(text)
            bl = ExpBL(doc=d, number=nums[0] if nums else "", po=find_po(text),
                       invoice_refs=[m.group(1) for m in core.INV_NO_RE.finditer(text)],
                       net=core.find_weight(text, core.TOTAL_NET_RE, dec),
                       gross=core.find_weight(text, core.TOTAL_GROSS_RE, dec),
                       packages=core.find_packages(text))
            bl.lines = parse_bl_lines(d, dec)
            bls.append(bl)
    # BL scindé sur plusieurs documents avec le même numéro : on fusionne (même BL)
    merged: "OrderedDict[str, ExpBL]" = OrderedDict()
    for bl in bls:
        key = bl.number or f"__{id(bl)}"
        if key in merged:
            tgt = merged[key]
            tgt.lines.extend(bl.lines)
            for attr in ("net", "gross", "packages"):
                a, b = getattr(tgt, attr), getattr(bl, attr)
                if a is None:
                    setattr(tgt, attr, b)
                elif b is not None and abs(a - b) > 1e-6:
                    notes.append(f"BL {bl.number} : {attr} différent entre deux documents ({a} / {b}) — à vérifier.")
            tgt.po = tgt.po or bl.po
            tgt.invoice_refs += bl.invoice_refs
        else:
            merged[key] = bl
    return invoices, list(merged.values())


# =============================================================================
# RAPPROCHEMENT BL ↔ FACTURES ET CONSTRUCTION DES LIGNES
# =============================================================================

def link_bl_invoices(invoices, bls, docs, notes):
    order = {id(d): i for i, d in enumerate(docs)}
    for bl in bls:
        linked = []
        for inv in invoices:
            by_bl = bl.number and any(core.compact_ref(r) == core.compact_ref(bl.number) for r in inv.bl_refs) \
                or (bl.number and len(bl.number) >= 5 and bl.number in inv.doc.text)
            by_inv = inv.number and any(core.compact_ref(r) == core.compact_ref(inv.number) for r in bl.invoice_refs)
            if by_bl or by_inv:
                linked.append(inv)
        if not linked and invoices:
            # repli : recouvrement des références
            refs = Counter(core.compact_ref(l["ref"]) for l in bl.lines)
            scored = [(sum(1 for l in inv.lines if core.compact_ref(l["ref"]) in refs), inv) for inv in invoices]
            best = max(scored, key=lambda t: t[0])
            if best[0] > 0:
                linked = [best[1]]
                notes.append(f"BL {bl.number or '?'} : facture associée par les références ({best[1].number}) — "
                             "aucun numéro croisé dans les documents.")
        if not linked and invoices:
            pos = order[id(bl.doc)]
            near = min(invoices, key=lambda inv: abs(order[id(inv.doc)] - pos))
            linked = [near]
            notes.append(f"BL {bl.number or '?'} : facture associée par proximité dans le dossier ({near.number}) "
                         "— correspondance incertaine, à vérifier.")
        bl.invoices = linked


def _match_invoice_line(bl_line, invoices):
    key = core.compact_ref(bl_line["ref"])
    for exact_qty in (True, False):
        for inv in invoices:
            for l in inv.lines:
                if l["_used"] or core.compact_ref(l["ref"]) != key:
                    continue
                if exact_qty and (l["qty"] is None or bl_line["qty"] is None or float(l["qty"]) != float(bl_line["qty"])):
                    continue
                l["_used"] = True
                return inv, l, exact_qty
    return None, None, False


def build_export_rows(invoices, bls, notes):
    rows = []

    def add(ref, desc, qty, unit, bl_no, inv_no, amount, po, group):
        rows.append({"Reference": ref, "Description": desc, "Quantity": qty, "Unit": unit,
                     "BL Number (Bon de Livraison)": bl_no, "Invoice Number (Facture)": inv_no,
                     "Total Net Weight of the BL": None, "Total Gross Weight of the BL": None,
                     "Number of Packages (Nombre de Colis)": None, "Prix unitaire total (Hors-taxe)": amount,
                     "Customer PO": po, "_group": group})

    bl_meta = OrderedDict()
    for g, bl in enumerate(bls):
        bl_meta[g] = bl
        for bl_line in bl.lines:
            inv, inv_line, exact = _match_invoice_line(bl_line, bl.invoices)
            if inv_line is None:
                inv_no = bl.invoices[0].number if len(bl.invoices) == 1 else ""
                notes.append(f"BL {bl.number or '?'} : réf {bl_line['ref']} (qté {bl_line['qty']}) sans ligne de facture "
                             "correspondante — description BL, montant HT vide.")
                add(bl_line["ref"], bl_line["desc"], bl_line["qty"], bl_line["unit"], bl.number, inv_no, None,
                    bl_line["po"] or bl.po or (bl.invoices[0].po if len(bl.invoices) == 1 else ""), g)
                continue
            if not exact:
                notes.append(f"BL {bl.number or '?'} : réf {bl_line['ref']} rapprochée de la facture {inv.number} "
                             f"avec une quantité différente (BL {bl_line['qty']} / facture {inv_line['qty']}).")
            if inv_line["amount"] is None:
                notes.append(f"Facture {inv.number} : réf {inv_line['ref']} sans montant HT identifiable.")
            add(bl_line["ref"], inv_line["desc"] or bl_line["desc"], bl_line["qty"], bl_line["unit"] or inv_line["unit"],
                bl.number, inv.number, inv_line["amount"], bl_line["po"] or inv_line["po"] or bl.po or inv.po, g)

    # Lignes de facture non couvertes par un BL
    g_next = len(bls)
    for inv in invoices:
        rest = [l for l in inv.lines if not l["_used"]]
        if not rest:
            continue
        linked_bls = [g for g, bl in bl_meta.items() if inv in bl.invoices]
        if linked_bls:
            g = linked_bls[0]
            bl = bl_meta[g]
            for l in rest:
                l["_used"] = True
                notes.append(f"Facture {inv.number} : réf {l['ref']} (qté {l['qty']}) absente du BL {bl.number or '?'} "
                             "— ligne ajoutée au BL associé, à vérifier.")
                add(l["ref"], l["desc"], l["qty"], l["unit"], bl.number, inv.number, l["amount"], l["po"] or bl.po or inv.po, g)
        else:
            bl_no = inv.bl_refs[0] if len(inv.bl_refs) == 1 else ""
            if bls and not bl_no:
                notes.append(f"Facture {inv.number} : aucun BL associé — lignes reprises sans numéro de BL.")
            fake = ExpBL(doc=inv.doc, number=bl_no, po=inv.po, net=inv.net, gross=inv.gross, packages=inv.packages)
            fake.invoices = [inv]
            bl_meta[g_next] = fake
            for l in rest:
                l["_used"] = True
                add(l["ref"], l["desc"], l["qty"], l["unit"], bl_no, inv.number, l["amount"], l["po"] or inv.po, g_next)
            g_next += 1

    # Regroupement par BL (lignes consécutives) en conservant l'ordre d'apparition
    rows.sort(key=lambda r: list(bl_meta).index(r["_group"]))
    seen = set()
    for r in rows:
        g = r["_group"]
        if g not in seen:
            seen.add(g)
            bl = bl_meta[g]
            r["Total Net Weight of the BL"] = bl.net
            r["Total Gross Weight of the BL"] = bl.gross
            r["Number of Packages (Nombre de Colis)"] = bl.packages
            if bl.net is None or bl.gross is None or bl.packages is None:
                miss = [n for n, v in (("poids net", bl.net), ("poids brut", bl.gross), ("colis", bl.packages)) if v is None]
                notes.append(f"BL {bl.number or '?'} : {', '.join(miss)} non trouvé(s) dans les documents.")
            # Écart de poids entre BL et facture
            for inv in bl.invoices:
                for lab, a, b in (("poids brut", bl.gross, inv.gross), ("poids net", bl.net, inv.net)):
                    if a is not None and b is not None and abs(a - b) > 0.01 and len(bl.invoices) == 1 \
                            and sum(1 for x in bl_meta.values() if inv in x.invoices) == 1:
                        notes.append(f"BL {bl.number} : {lab} BL {a} ≠ facture {inv.number} {b} — valeur du BL retenue.")
    df = pd.DataFrame(rows, columns=BLREF_COLUMNS + ["_group"])
    for c in ("Quantity", "Total Net Weight of the BL", "Total Gross Weight of the BL",
              "Number of Packages (Nombre de Colis)", "Prix unitaire total (Hors-taxe)"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    currencies = {g: (bl.invoices[0].currency if bl.invoices else "") for g, bl in bl_meta.items()}
    df["_currency"] = df["_group"].map(currencies).fillna("")
    return df


def build_factures_df(invoices, notes) -> pd.DataFrame:
    rows = []
    for inv in invoices:
        label = inv.number or f"({inv.doc.file} p.{inv.doc.page_label})"
        for fld, val in (("numéro", inv.number), ("nom client", inv.name), ("date", inv.date),
                         ("devise", inv.currency), ("montant HT", inv.total_ht)):
            if val in (None, ""):
                notes.append(f"Facture {label} : {fld} non trouvé(e) — cellule vide, à compléter avant import.")
        rows.append({"Numéro de facture": inv.number, "Nom du fournisseur ou client": inv.name,
                     "Date de la facture": inv.date, "Devise": inv.currency, "Montant en devise": inv.total_ht})
        # contrôle : somme des lignes vs total HT
        s = sum(l["amount"] for l in inv.lines + inv.fees if l["amount"] is not None)
        if inv.total_ht is not None and s and abs(s - inv.total_ht) > 0.02 and inv.total_src != "Somme des lignes":
            notes.append(f"Facture {label} : somme des lignes {s:.2f} ≠ total HT {inv.total_ht:.2f} "
                         f"(écart {inv.total_ht - s:+.2f}) — lignes manquantes, remises ou frais ?")
        if inv.fees:
            notes.append(f"Facture {label} : {len(inv.fees)} ligne(s) de frais non reprise(s) comme références "
                         f"({', '.join(core.clean_cell(f['desc'])[:25] for f in inv.fees)}).")
    return pd.DataFrame(rows, columns=FACTURES_COLUMNS)


def run_export_pipeline(pdf_files, ocr, client_mode: str, progress) -> dict:
    log, notes = [], []
    progress.progress(0.02, text="Étape 1/3 : Lecture des PDF (OCR si nécessaire)...")
    pages = []
    for f_idx, f in enumerate(pdf_files):
        def cb(i, total, f_idx=f_idx, name=f.name):
            progress.progress(0.02 + 0.6 * (f_idx + i / max(total, 1)) / len(pdf_files),
                              text=f"Étape 1/3 : Lecture des PDF... ({name} p.{i}/{total})")
        pages.extend(core.read_pdf_pages(f.name, f.getvalue(), ocr, log, cb))
    log.append(f"📄 {len(pages)} pages lues ({sum(p.ocr for p in pages)} via OCR)")
    docs = core.segment_documents(pages)
    progress.progress(0.65, text="Étape 2/3 : Extraction des lignes et rapprochement BL ↔ factures...")
    invoices, bls = parse_export_documents(docs, client_mode, notes)
    if not invoices:
        notes.append("Aucune facture commerciale détectée dans les PDF.")
    if not bls:
        notes.append("Aucun BL / Packing List détecté : les lignes viennent des factures, sans poids ni colis de BL.")
    link_bl_invoices(invoices, bls, docs, notes)
    bl_df = build_export_rows(invoices, bls, notes)
    fact_df = build_factures_df(invoices, notes)
    docs_df = pd.DataFrame([{"Fichier": d.file, "Pages": d.page_label, "Type": d.kind,
                             "N° détecté": (core.find_invoice_number(d.text) if d.kind == "FACTURE"
                                            else ", ".join(find_bl_numbers(d.text)[:2])),
                             "OCR": "Oui" if any(p.ocr for p in d.pages) else "Non"} for d in docs])
    # contrôle des occurrences : lignes sources vs lignes Excel
    src = sum(len(bl.lines) for bl in bls)
    log.append(f"🔁 Lignes d'articles BL : {src} source(s) → {int(bl_df['_group'].isin(range(len(bls))).sum())} ligne(s) Excel")
    progress.progress(0.85, text="Étape 3/3 : Génération des 2 fichiers Excel...")
    return {"factures": fact_df, "blref": bl_df, "docs": docs_df, "log": log, "notes": list(dict.fromkeys(notes))}


# =============================================================================
# FICHIERS EXCEL
# =============================================================================

HDR_FILL = PatternFill("solid", start_color="1F3864", end_color="1F3864")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
THIN = Side(style="thin", color="A6A6A6")
BRD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BL_COLORS = ["DDEBF7", "E2EFDA", "FFF2CC", "EDE1F5", "DDF1F0", "FBE4D5", "EDEDED", "E7EEF9", "F2F7E8", "FDEDEC"]
KG_FMT = '#,##0.000 "KG"'


def _autowidth(ws, minimum=8, maximum=60):
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            v = c.value
            ln = 10 if isinstance(v, (date, datetime)) else len(str(v))
            widths[c.column_letter] = max(widths.get(c.column_letter, 0), ln)
    for col, w in widths.items():
        ws.column_dimensions[col].width = max(minimum, min(maximum, w + 3))


def _to_date(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, pd.Timestamp):
        return v.date()
    return parse_date_any(str(v))


def _num(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v == "":
        return None
    if isinstance(v, str):
        v = core.parse_number(v, ",") if re.search(r",\d{1,2}$", v.strip()) else core.parse_number(v, ".")
        if v is None:
            return None
    return float(v)


def _txt(v) -> str | None:
    s = core.clean_cell(v)
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s or None


def build_factures_xlsx(df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Factures"
    ws.append(FACTURES_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    r = 2
    for _, row in df.iterrows():
        vals = [_txt(row["Numéro de facture"]), _txt(row["Nom du fournisseur ou client"]),
                _to_date(row["Date de la facture"]), (_txt(row["Devise"]) or "").upper() or None,
                _num(row["Montant en devise"])]
        if not any(v is not None for v in vals):
            continue  # aucune ligne vide
        a = ws.cell(row=r, column=1, value=vals[0]); a.number_format = "@"
        b = ws.cell(row=r, column=2, value=vals[1]); b.number_format = "@"
        c = ws.cell(row=r, column=3, value=vals[2]); c.number_format = "DD/MM/YYYY"
        d = ws.cell(row=r, column=4, value=vals[3]); d.number_format = "@"
        e = ws.cell(row=r, column=5, value=None if vals[4] is None else round(vals[4], 2)); e.number_format = "0.00"
        r += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(r - 1, 1)}"
    _autowidth(ws)
    ws.column_dimensions["C"].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_bl_xlsx(df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BL References"
    ws.append(BLREF_COLUMNS)
    for c in ws[1]:
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BRD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    text_cols = {"Reference", "Description", "Unit", "BL Number (Bon de Livraison)", "Invoice Number (Facture)", "Customer PO"}
    fmt = {"Quantity": "0.###", "Total Net Weight of the BL": KG_FMT, "Total Gross Weight of the BL": KG_FMT,
           "Number of Packages (Nombre de Colis)": "0", "Prix unitaire total (Hors-taxe)": "0.00"}
    groups = list(dict.fromkeys(df["_group"].tolist())) if "_group" in df else []
    n = 0
    for i, (_, row) in enumerate(df.iterrows(), start=2):
        color = BL_COLORS[groups.index(row["_group"]) % len(BL_COLORS)] if groups else None
        for j, col in enumerate(BLREF_COLUMNS, start=1):
            v = row[col]
            if col in text_cols:
                v = _txt(v)
            else:
                v = _num(v)
                if v is not None and col in ("Quantity", "Number of Packages (Nombre de Colis)") and float(v).is_integer():
                    v = int(v)
                if v is not None and col == "Prix unitaire total (Hors-taxe)":
                    v = round(v, 2)
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = BODY, BRD
            c.number_format = "@" if col in text_cols else fmt[col]
            if color:
                c.fill = PatternFill("solid", start_color=color, end_color=color)
        n = i
    last = max(n, 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{max(n, 1)}"
    _autowidth(ws)
    ws.column_dimensions["B"].width = min(70, max(30, ws.column_dimensions["B"].width))

    # ---------------- BL SUMMARY
    wsum = wb.create_sheet("BL SUMMARY")
    wsum.append(SUMMARY_COLUMNS)
    for c in wsum[1]:
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BRD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    R = lambda col: f"'BL References'!${col}$2:${col}${last}"  # noqa: E731
    bl_order = list(dict.fromkeys(df["BL Number (Bon de Livraison)"].map(lambda v: _txt(v) or "").tolist()))
    for i, bl in enumerate(bl_order, start=2):
        sub = df[df["BL Number (Bon de Livraison)"].map(lambda v: _txt(v) or "") == bl]
        if bl:
            crit = f"{R('E')},$A{i}"
            vals = [bl, f"=SUMIFS({R('C')},{crit})", f"=SUMIFS({R('G')},{crit})", f"=SUMIFS({R('H')},{crit})",
                    f"=SUMIFS({R('I')},{crit})", f"=SUMIFS({R('J')},{crit})"]
        else:  # BL non identifié : valeurs calculées (pas de critère vide fiable)
            vals = [None] + [float(pd.to_numeric(sub[c], errors="coerce").sum()) for c in
                             ("Quantity", "Total Net Weight of the BL", "Total Gross Weight of the BL",
                              "Number of Packages (Nombre de Colis)", "Prix unitaire total (Hors-taxe)")]
        for j, v in enumerate(vals, start=1):
            c = wsum.cell(row=i, column=j, value=v)
            c.font, c.border = BODY, BRD
            c.number_format = ["@", "0.###", KG_FMT, KG_FMT, "0", "#,##0.00"][j - 1]
    s_last = max(len(bl_order) + 1, 2)
    wsum.freeze_panes = "A2"
    wsum.auto_filter.ref = f"A1:F{s_last}"
    for col, w in zip("ABCDEF", (24, 16, 20, 20, 20, 20)):
        wsum.column_dimensions[col].width = w

    # ---------------- GLOBAL TOTAL
    wg = wb.create_sheet("GLOBAL TOTAL")
    wg["A1"] = "TOTAL GLOBAL — TOUS LES BL"
    wg["A1"].font = Font(name="Arial", size=13, bold=True, color="1F3864")
    for j, h in enumerate(("Champ", "Valeur"), start=1):
        c = wg.cell(row=3, column=j, value=h)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BRD
    S = lambda col: f"'BL SUMMARY'!{col}2:{col}{s_last}"  # noqa: E731
    currencies = [c for c in dict.fromkeys(df.get("_currency", pd.Series(dtype=str)).fillna("").tolist())]
    rows = [("COLIS", f"=SUM({S('E')})", "0"), ("POIDS BRUT", f"=SUM({S('D')})", KG_FMT),
            ("POIDS NET", f"=SUM({S('C')})", KG_FMT)]
    if len([c for c in currencies if c]) <= 1:
        cur = next((c for c in currencies if c), "")
        rows.append(("VALEUR DEVISE", f"=SUM({S('F')})", f'#,##0.00 "{cur}"' if cur else "#,##0.00"))
    else:  # devises différentes : jamais additionnées entre elles
        for cur in currencies:
            tot = pd.to_numeric(df.loc[df["_currency"] == cur, "Prix unitaire total (Hors-taxe)"], errors="coerce").sum()
            rows.append((f"VALEUR DEVISE ({cur or 'devise ?'})", round(float(tot), 2),
                         f'#,##0.00 "{cur}"' if cur else "#,##0.00"))
    rows.append(("QUANTITE", f"=SUM({S('B')})", "#,##0.###"))
    for i, (lab, val, nf) in enumerate(rows, start=4):
        a = wg.cell(row=i, column=1, value=lab)
        b = wg.cell(row=i, column=2, value=val)
        a.font, a.border, b.font, b.border = BOLD, BRD, Font(name="Arial", size=11), BRD
        a.fill = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
        b.number_format = nf
        b.alignment = Alignment(horizontal="right")
    wg.column_dimensions["A"].width = 28
    wg.column_dimensions["B"].width = 26

    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =============================================================================
# VÉRIFICATION APRÈS CRÉATION (relecture des fichiers)
# =============================================================================

def verify_files(fact_bytes: bytes, bl_bytes: bytes) -> list[tuple[bool, str]]:
    checks = []
    wb = load_workbook(io.BytesIO(fact_bytes))
    checks.append((wb.sheetnames == ["Factures"], "FACTURES.xlsx : une seule feuille nommée « Factures »"))
    ws = wb["Factures"]
    hdr = [c.value for c in ws[1]]
    checks.append((hdr == FACTURES_COLUMNS and ws.max_column == 5, "FACTURES.xlsx : exactement 5 colonnes, en-têtes exacts"))
    data_rows = list(ws.iter_rows(min_row=2))
    checks.append((all(any(c.value is not None for c in r) for r in data_rows), "Aucune ligne vide"))
    checks.append((not ws.merged_cells.ranges, "Aucune cellule fusionnée"))
    checks.append((all(isinstance(r[0].value, str) or r[0].value is None for r in data_rows), "Colonne A : numéros en texte"))
    checks.append((all(r[2].value is None or (isinstance(r[2].value, (date, datetime)) and r[2].number_format == "DD/MM/YYYY")
                       for r in data_rows), "Colonne C : vraies dates Excel au format DD/MM/YYYY"))
    checks.append((all(r[3].value is None or re.fullmatch(r"[A-Z]{3}", str(r[3].value)) for r in data_rows),
                   "Colonne D : codes devise sur 3 lettres"))
    checks.append((all(r[4].value is None or (isinstance(r[4].value, (int, float)) and r[4].number_format == "0.00")
                       for r in data_rows), "Colonne E : montants numériques au format 0.00 (jamais du texte)"))
    complete = all(all(c.value is not None for c in r[:5]) for r in data_rows)
    checks.append((complete, "Toutes les cellules des factures sont renseignées" if complete
                   else "Certaines cellules de FACTURES sont vides (voir compte rendu) — l'import système les refusera"))
    wb2 = load_workbook(io.BytesIO(bl_bytes))
    checks.append((wb2.sheetnames == ["BL References", "BL SUMMARY", "GLOBAL TOTAL"],
                   "BL_REFERENCES_SUMMARY.xlsx : 3 feuilles dans l'ordre exact"))
    checks.append(([c.value for c in wb2["BL References"][1]] == BLREF_COLUMNS, "BL References : 11 colonnes exactes"))
    checks.append(([c.value for c in wb2["BL SUMMARY"][1]] == SUMMARY_COLUMNS, "BL SUMMARY : 6 colonnes exactes"))
    checks.append((not any(w.merged_cells.ranges for w in wb2.worksheets), "Aucune cellule fusionnée"))
    # poids/colis uniquement sur la 1re ligne de chaque BL
    ws = wb2["BL References"]
    seen, ok = set(), True
    for r in ws.iter_rows(min_row=2, values_only=True):
        bl = r[4]
        if bl in seen and any(v is not None for v in r[6:9]):
            ok = False
        seen.add(bl)
    checks.append((ok, "Poids et colis présents uniquement sur la première ligne de chaque BL"))
    return checks


def zip_files(files: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


def build_report(fact_df: pd.DataFrame, bl_df: pd.DataFrame, notes: list) -> dict:
    refs = bl_df["Reference"].map(lambda v: core.clean_cell(v)).replace("", pd.NA).dropna()
    per_cur = (bl_df.assign(v=pd.to_numeric(bl_df["Prix unitaire total (Hors-taxe)"], errors="coerce"))
               .groupby(bl_df.get("_currency", pd.Series([""] * len(bl_df))).fillna(""))["v"].sum().to_dict()
               if len(bl_df) else {})
    num = lambda c: float(pd.to_numeric(bl_df[c], errors="coerce").sum()) if len(bl_df) else 0.0  # noqa: E731
    repeated = int((refs.value_counts() > 1).sum())
    return {
        "factures": len(fact_df), "bl": bl_df["BL Number (Bon de Livraison)"].map(lambda v: core.clean_cell(v)).nunique() if len(bl_df) else 0,
        "lignes": len(bl_df), "refs_distinctes": int(refs.nunique()), "refs_repetees": repeated,
        "colis": num("Number of Packages (Nombre de Colis)"), "poids_brut": num("Total Gross Weight of the BL"),
        "poids_net": num("Total Net Weight of the BL"), "quantite": num("Quantity"), "valeur_par_devise": per_cur,
        "factures_par_devise": (fact_df.assign(v=pd.to_numeric(fact_df["Montant en devise"], errors="coerce"))
                                .groupby("Devise")["v"].sum().to_dict() if len(fact_df) else {}),
        "notes": notes,
    }


def report_text(rep: dict, checks: list) -> str:
    lines = ["COMPTE RENDU — EXPORT BL / FACTURES", "=" * 40,
             f"Factures extraites           : {rep['factures']}",
             f"BL uniques                   : {rep['bl']}",
             f"Lignes d'articles extraites  : {rep['lignes']}",
             f"Références distinctes        : {rep['refs_distinctes']}",
             f"Références répétées          : {rep['refs_repetees']} (chaque occurrence conservée sur sa propre ligne)",
             f"Total colis                  : {rep['colis']:g}",
             f"Poids brut total             : {rep['poids_brut']:.3f} KG",
             f"Poids net total              : {rep['poids_net']:.3f} KG",
             f"Quantité totale              : {rep['quantite']:g}"]
    for cur, v in rep["valeur_par_devise"].items():
        lines.append(f"Valeur totale HT ({cur or '?'})      : {v:.2f}")
    lines += ["", "CONTRÔLES DES FICHIERS"] + [f"[{'OK' if ok else '!!'}] {msg}" for ok, msg in checks]
    lines += ["", "ANOMALIES / POINTS À VÉRIFIER"] + ([f"- {n}" for n in rep["notes"]] or ["- Aucune"])
    return "\n".join(lines)


# =============================================================================
# INTERFACE STREAMLIT — MODULE EXPORT
# =============================================================================

def normalize_bl_df(df: pd.DataFrame) -> pd.DataFrame:
    """Après édition : groupes et devises recalculés à partir du numéro de BL."""
    df = df.copy()
    for c in BLREF_COLUMNS:
        if c not in df.columns:
            df[c] = None
    bl = df["BL Number (Bon de Livraison)"].map(lambda v: core.clean_cell(v))
    df["_group"] = bl.map({b: i for i, b in enumerate(dict.fromkeys(bl))})
    if "_currency" not in df.columns:
        df["_currency"] = ""
    cur = df["_currency"].map(lambda v: core.clean_cell(v))
    by_bl = {b: c for b, c in zip(bl, cur) if c}
    df["_currency"] = [c or by_bl.get(b, "") for b, c in zip(bl, cur)]
    order = list(dict.fromkeys(bl))
    df["_o"] = bl.map(order.index)
    df = df.sort_values("_o", kind="stable").drop(columns="_o")  # lignes d'un même BL consécutives
    return df.reset_index(drop=True)


def render_export(ocr):
    import streamlit as st
    import ui

    res = st.session_state.get("exp_result")
    if res:
        rep0 = build_report(res["factures"], res["blref"], res["notes"])
        stats = {"factures": rep0["factures"], "BL uniques": rep0["bl"], "lignes articles": rep0["lignes"],
                 "colis": f"{rep0['colis']:g}"}
    else:
        stats = {"fichiers Excel": 2, "FACTURES": "import système", "BL": "3 feuilles", "OCR": "intégré"}
    ui.hero_3d(stats, badge="TIMAR TANGER · EXPORT ALSTOM", title="Export BL & Factures",
               highlight="2 fichiers Excel",
               desc="BL / Packing Lists et factures commerciales lus automatiquement : FACTURES.xlsx prêt pour "
                    "« Créer une déclaration » et BL_REFERENCES_SUMMARY.xlsx — une ligne source = une ligne Excel.")

    c1, c2 = st.columns([1.6, 1], gap="medium")
    with c1:
        with st.container(key="card_upl_exp"):
            ui.card_title("📂", "Dossier « extrait » — tous les PDF export",
                          "Packing Lists / BL et Commercial Invoices / Factures (un ou plusieurs PDF)",
                          ["Packing List", "Bon de Livraison", "Commercial Invoice", "Facture", "Scans OCR"])
            files = st.file_uploader("PDF export", type=["pdf"], accept_multiple_files=True, key="exp_upl")
    with c2:
        with st.container(key="card_upl_opt"):
            ui.card_title("⚙️", "Options FACTURES.xlsx", "Colonne B : « Nom du fournisseur ou client »")
            mode = st.radio("Nom à reporter en colonne B", ["Client facturé (Sold to / Facturé à)",
                                                              "Fournisseur (émetteur de la facture)"],
                            label_visibility="collapsed", key="exp_client_mode")
            st.caption("Par défaut : le client facturé, comme demandé dans le cahier des charges.")

    ui.steps_strip(2 if res else (1 if files else 0), [
        ("📥", "Importer", "BL / Packing Lists + Factures"),
        ("🧠", "Extraire & rapprocher", "Lignes, poids BL, montants HT"),
        ("📊", "Générer", "FACTURES + BL_REFERENCES_SUMMARY")])

    go = st.button("🚀  Générer FACTURES.xlsx + BL_REFERENCES_SUMMARY.xlsx", type="primary",
                   use_container_width=True, disabled=not files, key="exp_go")
    if go:
        progress = st.progress(0.0, text="Initialisation...")
        try:
            r = run_export_pipeline(files, ocr, "client" if mode.startswith("Client") else "supplier", progress)
        except Exception as exc:
            progress.empty()
            st.error(f"Erreur pendant le traitement : {exc}")
            st.exception(exc)
            return
        st.session_state["exp_result"] = r
        st.session_state["exp_ver"] = st.session_state.get("exp_ver", 0) + 1
        progress.progress(1.0, text="✅ Les 2 fichiers Excel sont prêts.")
        st.rerun()

    if not res:
        return
    ver = st.session_state.get("exp_ver", 0)

    t_dash, t_fact, t_bl, t_rep, t_docs = st.tabs(["📊  Résumé & téléchargements", "🧾  FACTURES (modifiable)",
                                                   "📦  BL References (modifiable)", "✅  Compte rendu & contrôles",
                                                   "🗂️  Documents & journal"])
    with t_fact:
        st.caption("Une ligne par facture. Corrigez si besoin : le fichier est régénéré et revérifié automatiquement.")
        fdf = res["factures"].copy()
        fdf["Date de la facture"] = pd.to_datetime(fdf["Date de la facture"], errors="coerce")
        fact_edit = st.data_editor(fdf, num_rows="dynamic", use_container_width=True, key=f"exp_fact_{ver}",
                                   column_config={
                                       "Date de la facture": st.column_config.DateColumn(format="DD/MM/YYYY"),
                                       "Montant en devise": st.column_config.NumberColumn(format="%.2f"),
                                       "Numéro de facture": st.column_config.TextColumn()})
    with t_bl:
        st.caption("Une ligne source = une ligne Excel. Les références répétées restent sur des lignes séparées ; "
                   "poids et colis uniquement sur la 1re ligne de chaque BL.")
        bl_edit = st.data_editor(res["blref"], num_rows="dynamic", use_container_width=True, height=460,
                                 key=f"exp_bl_{ver}",
                                 column_config={"_group": None, "_currency": None,
                                                "Prix unitaire total (Hors-taxe)": st.column_config.NumberColumn(format="%.2f")})

    fact_final = fact_edit.copy()
    fact_final["Date de la facture"] = fact_final["Date de la facture"].map(_to_date)
    bl_final = normalize_bl_df(bl_edit)
    for c in ("Quantity", "Total Net Weight of the BL", "Total Gross Weight of the BL",
              "Number of Packages (Nombre de Colis)", "Prix unitaire total (Hors-taxe)"):
        bl_final[c] = pd.to_numeric(bl_final[c], errors="coerce")
    fact_bytes = build_factures_xlsx(fact_final)
    bl_bytes = build_bl_xlsx(bl_final)
    checks = verify_files(fact_bytes, bl_bytes)
    rep = build_report(fact_final, bl_final, res["notes"])
    report = report_text(rep, checks)

    with t_dash:
        ui.kpi_cards([
            {"icon": "🧾", "label": "Factures", "value": rep["factures"], "sub": "1 ligne par facture", "tone": "blue"},
            {"icon": "🚚", "label": "BL uniques", "value": rep["bl"], "sub": "bons de livraison", "tone": "cyan"},
            {"icon": "📋", "label": "Lignes articles", "value": rep["lignes"], "sub": "1 ligne source = 1 ligne", "tone": "gold"},
            {"icon": "🔖", "label": "Réf. distinctes", "value": rep["refs_distinctes"], "sub": "codes uniques", "tone": "blue"},
            {"icon": "🔁", "label": "Réf. répétées", "value": rep["refs_repetees"], "sub": "conservées séparément",
             "tone": "good"},
            {"icon": "🔢", "label": "Quantité", "value": rep["quantite"], "decimals": 0 if float(rep["quantite"]).is_integer() else 2,
             "sub": "total des lignes", "tone": "cyan"},
        ])
        cards = [{"icon": "📦", "label": "Colis", "value": rep["colis"], "sub": "BL uniques", "tone": "gold"},
                 {"icon": "⚖️", "label": "Poids brut", "value": round(rep["poids_brut"], 3), "decimals": 3, "sub": "KG", "tone": "blue"},
                 {"icon": "🪶", "label": "Poids net", "value": round(rep["poids_net"], 3), "decimals": 3, "sub": "KG", "tone": "cyan"}]
        cards += [{"icon": "💶", "label": f"Valeur HT {cur or '(devise ?)'}", "value": round(float(v), 2), "decimals": 2,
                   "sub": cur, "tone": "good"} for cur, v in rep["valeur_par_devise"].items()]
        ui.kpi_cards(cards[:6])

        failed = [m for ok, m in checks if not ok]
        if failed:
            ui.status_banner("ko", "🔴", "Contrôle des fichiers : à corriger", " · ".join(failed))
        elif rep["notes"]:
            ui.status_banner("warn", "🟠", f"Fichiers conformes — {len(rep['notes'])} point(s) à vérifier",
                             "Voir l'onglet « Compte rendu & contrôles ».")
        else:
            ui.status_banner("ok", "🟢", "Fichiers conformes — aucune anomalie",
                             "Les deux fichiers ont été relus et contrôlés après création.")

        d1, d2, d3 = st.columns(3)
        with d1:
            st.download_button("📥  FACTURES.xlsx", data=fact_bytes, file_name="FACTURES.xlsx", type="primary",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               use_container_width=True, key="dl_fact")
        with d2:
            st.download_button("📥  BL_REFERENCES_SUMMARY.xlsx", data=bl_bytes, file_name="BL_REFERENCES_SUMMARY.xlsx",
                               type="primary", use_container_width=True, key="dl_bl",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        with d3:
            st.download_button("🗜️  Les 2 fichiers + compte rendu (.zip)", key="dl_zip", type="primary",
                               data=zip_files({"FACTURES.xlsx": fact_bytes, "BL_REFERENCES_SUMMARY.xlsx": bl_bytes,
                                               "COMPTE_RENDU.txt": report.encode("utf-8")}),
                               file_name="EXPORT_BL_FACTURES.zip", mime="application/zip", use_container_width=True)

        ui.section("BL SUMMARY", "🚚")
        b = bl_final.copy()
        b["BL"] = b["BL Number (Bon de Livraison)"].map(core.clean_cell)
        summ = (b.groupby("BL", sort=False)
                .agg(**{"Total Quantity": ("Quantity", lambda s: pd.to_numeric(s, errors="coerce").sum()),
                        "Total Net Weight": ("Total Net Weight of the BL", lambda s: pd.to_numeric(s, errors="coerce").sum()),
                        "Total Gross Weight": ("Total Gross Weight of the BL", lambda s: pd.to_numeric(s, errors="coerce").sum()),
                        "Number of Packages": ("Number of Packages (Nombre de Colis)", lambda s: pd.to_numeric(s, errors="coerce").sum()),
                        "Total Valeur HT": ("Prix unitaire total (Hors-taxe)", lambda s: pd.to_numeric(s, errors="coerce").sum()),
                        "Devise": ("_currency", "first")}).reset_index().rename(columns={"BL": "BL Number"}))
        st.dataframe(summ, use_container_width=True, hide_index=True,
                     column_config={"Total Net Weight": st.column_config.NumberColumn(format="%.3f KG"),
                                    "Total Gross Weight": st.column_config.NumberColumn(format="%.3f KG"),
                                    "Total Valeur HT": st.column_config.NumberColumn(format="%.2f")})
        ui.section("Aperçu BL References", "📦")
        view = bl_final.drop(columns=["_group", "_currency"]).map(  # cellules vides hors 1re ligne du BL
            lambda v: "" if v is None or (isinstance(v, float) and pd.isna(v)) else (f"{v:g}" if isinstance(v, float) else str(v)))
        st.dataframe(view, use_container_width=True, hide_index=True, height=320)

    with t_rep:
        ui.section("Contrôles des fichiers (relus après création)", "✅")
        for ok, msg in checks:
            st.markdown(f"{'✅' if ok else '❌'} {msg}")
        ui.section("Anomalies / points à vérifier", "⚠️")
        if rep["notes"]:
            for n in rep["notes"]:
                st.markdown(f"- {n}")
        else:
            st.markdown("Aucune anomalie détectée.")
        ui.section("Compte rendu complet", "📝")
        st.code(report, language=None)

    with t_docs:
        ui.section("Documents détectés", "🗂️")
        st.dataframe(res["docs"], use_container_width=True, hide_index=True)
        ui.section("Journal", "📝")
        for m in res["log"]:
            st.write(m)
